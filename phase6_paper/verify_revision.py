#!/usr/bin/env python3
"""Verify all reviewer corrections landed in PERSIST_Revised_Corrected.docx."""
import re
import docx
from docx.enum.text import WD_COLOR_INDEX

P = ("/projects/sandbox/Assesment_of_Ecological_resilience_Yangtze/"
     "phase6_paper/PERSIST_Revised_Corrected.docx")
d = docx.Document(P)
txt = "\n".join(p.text for p in d.paragraphs)

n_hl_para = sum(1 for p in d.paragraphs
                for r in p.runs if r.font.highlight_color == WD_COLOR_INDEX.YELLOW)
n_hl_cell = sum(1 for t in d.tables for row in t.rows for c in row.cells
                for pp in c.paragraphs for r in pp.runs
                if r.font.highlight_color == WD_COLOR_INDEX.YELLOW)
print(f"YELLOW runs: {n_hl_para} in paragraphs, {n_hl_cell} in table cells")
print(f"paragraphs={len(d.paragraphs)}  tables={len(d.tables)}")

print("\n=== EQUATIONS (LaTeX present?) ===")
eqs = []
for p in d.paragraphs:
    m = re.search(r"\((\d+[abc]?)\)\s*$", p.text.strip())
    if m and ("\\" in p.text or "=" in p.text):
        eqs.append(m.group(1))
print(f"  found {len(eqs)}: {eqs}")
missing_latex = [p.text.strip() for p in d.paragraphs
                 if re.fullmatch(r"\(\d+[abc]?\)", p.text.strip())]
print(f"  empty equation slots remaining: {missing_latex if missing_latex else 'NONE'}")

CHECKS = [
    ("C1 dimension 5+22+2+3=32", "5+22+2+3=32" in txt and "twenty-two channels" in txt),
    ("C2 Table 5 present", any("Table 5." in p.text for p in d.paragraphs)),
    ("C3 ridge stride-H in Table 3",
     any("Ridge stride-H R2" in c.text for t in d.tables for c in t.rows[0].cells)),
    ("C4 kNDVI reported", "kNDVI R2 (secondary target)" in
     "\n".join(c.text for t in d.tables for row in t.rows for c in row.cells)),
    ("C5 significance tests", "Kruskal-Wallis" in txt and "Holm" in txt),
    ("C6 Fig5 (a)/(b) caption", "(a) Change in test RMSE" in txt and "(b) test RMSE" in txt),
    ("C7 Table 1 rebuilt", "Benchmarked in this study" in
     "\n".join(c.text for t in d.tables for row in t.rows for c in row.cells)),
    ("C8 40% sourced", "40% of China" in txt and "[2], [38]" in txt),
    ("C9 karst ref versioned", "WOKAM), 2017 edition" in txt),
    ("C10 NTL extended cited", "[58]" in txt and "1992-2018 record" in txt),
    ("C11 no bare-URL refs", not any(re.fullmatch(r"\[\d+\]\s*https?://\S+", p.text.strip())
                                    for p in d.paragraphs)),
    ("C12 H=12 within value", "-0.7040" in txt),
    ("C13 0.82 at audit", "fixed before model development" in txt
     and "stated here where the audit is introduced" in txt),
    ("C14 edge cut corrected", "29.9% and 31.3%" in txt and "8.0% of edges" in txt),
    ("C15 update factor corrected", "137 in expectation" in txt or "136 in expectation" in txt),
    ("C16 decoder init wording", "initialisation toward the seasonal-naive" in txt
     and "rather than pre-training" in txt),
    ("C17 ridge basis explained", "not computed on the same basis" in txt
     and "0.8884" in txt),
    ("C18 sigma per variable", "single global standard deviation of variable" in txt
     or "one global scale per variable" in txt),
    ("C19 DataV de-linked", "[57]" in txt and "DataV.GeoAtlas" in txt),
    ("C20 total loss eq added", "(15c)" in txt),
]
print("\n=== CHECKS ===")
bad = 0
for lab, ok in CHECKS:
    print(f"  [{'PASS' if ok else 'FAIL'}] {lab}")
    bad += (not ok)

print("\n=== TABLE 3 ===")
for row in d.tables[2].rows:
    print("  ", [c.text.strip() for c in row.cells])

print("\n=== TABLE 12 (significance) ===")
for row in d.tables[11].rows:
    print("  ", [c.text.strip()[:22] for c in row.cells])

print(f"\n{'ALL CHECKS PASSED' if bad == 0 else str(bad) + ' CHECK(S) FAILED'}")
