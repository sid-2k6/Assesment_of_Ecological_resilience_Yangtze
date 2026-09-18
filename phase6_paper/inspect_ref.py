#!/usr/bin/env python3
"""Dump the structure + formatting of the reference FOWT-ARISE document."""
import docx
from docx.shared import Pt

P = ("/projects/sandbox/Assesment_of_Ecological_resilience_Yangtze/"
     "Outputs_v3/Sample_document_for_reference/FOWT-ARISE.docx")
d = docx.Document(P)

sec = d.sections[0]
print("=== PAGE SETUP ===")
print(f"  page   {sec.page_width.inches:.2f} x {sec.page_height.inches:.2f} in")
print(f"  margins L{sec.left_margin.inches:.2f} R{sec.right_margin.inches:.2f} "
      f"T{sec.top_margin.inches:.2f} B{sec.bottom_margin.inches:.2f}")

st = d.styles["Normal"]
print(f"\n=== NORMAL STYLE ===\n  font {st.font.name}  size "
      f"{st.font.size.pt if st.font.size else None}")
pf = st.paragraph_format
print(f"  align {pf.alignment}  line_spacing {pf.line_spacing}  "
      f"space_after {pf.space_after}")

print(f"\n=== BODY: {len(d.paragraphs)} paragraphs, {len(d.tables)} tables ===\n")
for i, p in enumerate(d.paragraphs):
    t = p.text.strip()
    has_img = "graphic" in p._p.xml
    if not t and not has_img:
        continue
    runs = [r for r in p.runs if r.text.strip()]
    f = runs[0].font if runs else None
    fname = f.name if f else None
    fsize = f.size.pt if (f and f.size) else None
    bold = f.bold if f else None
    ital = f.italic if f else None
    al = p.alignment
    tag = "[IMG]" if has_img else ""
    print(f"{i:>4} | {p.style.name:<22} | {str(fname):<14} {str(fsize):<6} "
          f"b={str(bold):<5} i={str(ital):<5} al={str(al):<22} {tag} "
          f"{t[:95]}")

print("\n=== TABLES ===")
for ti, tb in enumerate(d.tables):
    print(f"\n--- Table {ti}: style={tb.style.name if tb.style else None} "
          f"rows={len(tb.rows)} cols={len(tb.columns)}")
    for ri, row in enumerate(tb.rows[:4]):
        cells = [c.text.strip().replace("\n", " / ")[:34] for c in row.cells]
        print(f"    r{ri}: {cells}")
    if len(tb.rows) > 4:
        print(f"    ... {len(tb.rows)-4} more rows")
    c0 = tb.rows[0].cells[0]
    if c0.paragraphs and c0.paragraphs[0].runs:
        rf = c0.paragraphs[0].runs[0].font
        print(f"    header font: {rf.name} {rf.size.pt if rf.size else None} "
              f"bold={rf.bold}")
    if len(tb.rows) > 1:
        c1 = tb.rows[1].cells[0]
        if c1.paragraphs and c1.paragraphs[0].runs:
            rf = c1.paragraphs[0].runs[0].font
            print(f"    body   font: {rf.name} {rf.size.pt if rf.size else None} "
                  f"bold={rf.bold}")
