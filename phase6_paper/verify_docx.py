#!/usr/bin/env python3
"""Verify the generated manuscript against the reference format."""
import re
import docx
from docx.shared import Emu

P = ("/projects/sandbox/Assesment_of_Ecological_resilience_Yangtze/"
     "phase6_paper/PERSIST_Manuscript.docx")
d = docx.Document(P)
s = d.sections[0]
print("=== PAGE SETUP ===")
print(f"  {s.page_width.inches:.2f} x {s.page_height.inches:.2f} in, margins "
      f"{s.left_margin.inches:.2f}/{s.right_margin.inches:.2f}/"
      f"{s.top_margin.inches:.2f}/{s.bottom_margin.inches:.2f}")
st = d.styles["Normal"]
print(f"  Normal: {st.font.name} {st.font.size.pt}pt")

imgs, eqs = [], []
for i, p in enumerate(d.paragraphs):
    if "graphic" in p._p.xml:
        for cx, cy in re.findall(r'<wp:extent cx="(\d+)" cy="(\d+)"', p._p.xml):
            imgs.append((i, Emu(int(cx)).inches, Emu(int(cy)).inches))
    m = re.search(r"\((\d+)\)\s*$", p.text.strip())
    if m and len(p.text) > 25:
        eqs.append(int(m.group(1)))

print(f"\n=== IMAGES: {len(imgs)} ===")
for i, w, h in imgs:
    print(f"  para {i:>3}: {w:.2f} x {h:.2f} in")

print(f"\n=== EQUATIONS: {len(eqs)} -> {eqs}")
gaps = [n for n in range(1, max(eqs) + 1) if n not in eqs] if eqs else []
print(f"  numbering gaps: {gaps if gaps else 'none'}")

print(f"\n=== HEADINGS ===")
for p in d.paragraphs:
    t = p.text.strip()
    if not t:
        continue
    runs = [r for r in p.runs if r.text.strip()]
    if runs and runs[0].bold and re.match(r"^(\d+\.|\d+\.\d+|Abstract|References|Keywords)", t):
        print(f"  {t[:80]}")

print(f"\n=== TABLES: {len(d.tables)} ===")
for ti, tb in enumerate(d.tables, 1):
    hdr = [c.text.strip()[:22] for c in tb.rows[0].cells]
    fs = tb.rows[0].cells[0].paragraphs[0].runs[0].font.size
    print(f"  Table {ti}: {len(tb.rows)-1} data rows x {len(tb.columns)} cols "
          f"@ {fs.pt if fs else '?'}pt  style={tb.style.name}")
    print(f"     {hdr}")

caps = [p.text.strip() for p in d.paragraphs
        if p.text.strip().startswith(("Figure ", "Table "))]
print(f"\n=== CAPTIONS: {len(caps)} ===")
for c in caps:
    print(f"  {c[:95]}")

words = sum(len(p.text.split()) for p in d.paragraphs)
print(f"\n=== WORD COUNT (body paragraphs): {words:,} ===")
