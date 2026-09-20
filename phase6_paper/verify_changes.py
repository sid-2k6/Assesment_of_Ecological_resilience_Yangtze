#!/usr/bin/env python3
"""Confirm the six reviewer fixes landed with correct revision mark-up."""
import docx
from docx.enum.text import WD_COLOR_INDEX

P = ("/projects/sandbox/Assesment_of_Ecological_resilience_Yangtze/"
     "phase6_paper/PERSIST_Manuscript.docx")
d = docx.Document(P)
NAME = {WD_COLOR_INDEX.RED: "REMOVED(red+strike)",
        WD_COLOR_INDEX.BRIGHT_GREEN: "ADDED(green)",
        WD_COLOR_INDEX.YELLOW: "REPLACED(yellow)"}

tally = {"REMOVED(red+strike)": 0, "ADDED(green)": 0, "REPLACED(yellow)": 0}
print("=" * 78)
print("MARKED RUNS IN BODY PARAGRAPHS")
print("=" * 78)
for i, p in enumerate(d.paragraphs):
    for r in p.runs:
        hc = r.font.highlight_color
        if hc in NAME:
            k = NAME[hc]
            tally[k] += 1
            strike = " [strike]" if r.font.strike else ""
            print(f"\npara {i:>3} | {k}{strike}")
            print(f"  {r.text[:300]}")

print("\n" + "=" * 78)
print("MARKED CELLS IN TABLES")
print("=" * 78)
for ti, tb in enumerate(d.tables, 1):
    for ri, row in enumerate(tb.rows):
        marks = []
        for ci, c in enumerate(row.cells):
            for pp in c.paragraphs:
                for r in pp.runs:
                    if r.font.highlight_color in NAME:
                        marks.append((ci, NAME[r.font.highlight_color], r.text))
        if marks:
            kinds = {m[1] for m in marks}
            for k in kinds:
                tally[k] += 1
            print(f"\nTable {ti} row {ri}: {sorted(kinds)}")
            print(f"  {[c.text.strip()[:26] for c in row.cells]}")

print("\n" + "=" * 78)
print("TALLY:", tally)
print("=" * 78)

# targeted assertions
txt = "\n".join(p.text for p in d.paragraphs)
checks = [
    ("Fix 1 peak claim removed",
     "it is the horizon at which genuine temporal skill" in txt
     and "most favourable trade-off rather than as the maximum" in txt),
    ("Fix 2 e-index clarified",
     "first month of the target window" in txt and "half-open interval" in txt),
    ("Fix 3 H=12 within value present",
     "-0.7040" in txt or "-0.704" in txt),
    ("Fix 4 0.82 pre-specified at audit",
     "pre-specified acceptance threshold" in txt),
    ("Fix 5 per-variable sigma",
     "separately for each raw variable" in txt),
    ("Fix 6 ridge basis explained",
     "are not computed on the same basis" in txt
     and "0.8884" in txt and "single-target basis" in txt),
]
print("\nASSERTIONS")
for label, ok in checks:
    print(f"  [{'PASS' if ok else 'FAIL'}] {label}")

t10 = d.tables[9]
print(f"\nTable 10 rows = {len(t10.rows)-1} (was 8, expect 10)")
for row in t10.rows[-3:]:
    print("  ", [c.text.strip()[:34] for c in row.cells][:6])
t3 = d.tables[2]
print("\nTable 3 H=12 row:", [c.text.strip() for c in t3.rows[-1].cells])
