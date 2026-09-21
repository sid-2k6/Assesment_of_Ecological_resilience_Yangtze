#!/usr/bin/env python3
"""Apply every reviewer correction to 'PERSIST 21.09.docx'.

All insertions and edits are highlighted YELLOW. Equations are inserted as
LaTeX so they can be selected in Word and converted via Equation > LaTeX.
"""
import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph

REPO = Path("/projects/sandbox/Assesment_of_Ecological_resilience_Yangtze")
SRC = REPO / "PERSIST 21.09.docx"
DST = REPO / "phase6_paper" / "PERSIST_Revised_Corrected.docx"
OUT = REPO / "Outputs_v3" / "outputs_v3"

M = json.load(open(OUT / "PERSIST" / "metrics.json"))
RIDGE = pd.read_csv(REPO / "phase5_proposed_model" / "horizon_scan.csv")
WSCAN = pd.read_csv(REPO / "phase5_proposed_model" / "within_scan.csv")
SIG = json.load(open("/projects/sandbox/yreb_resilience/sig_tests.json"))

doc = Document(str(SRC))
P = doc.paragraphs
Y = WD_COLOR_INDEX.YELLOW
LOG = []


def _style(r, fs=12, hl=True, italic=False):
    r.font.name = "Times New Roman"
    r.font.size = Pt(fs)
    r.font.italic = italic
    if hl:
        r.font.highlight_color = Y
    return r


def rewrite(idx, segments, fs=12):
    """segments = [(text, highlight_bool)]"""
    p = P[idx]
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    for text, hl in segments:
        _style(p.add_run(text), fs=fs, hl=hl)
    LOG.append(f"rewrote para {idx}")
    return p


def set_eq(idx, latex, num):
    p = P[idx]
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    _style(p.add_run(latex), fs=11, hl=True)
    _style(p.add_run(f"\t({num})"), fs=12, hl=False)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    LOG.append(f"eq ({num}) -> para {idx}")
    return p


def para_after(ref_p):
    el = OxmlElement("w:p")
    ref_p._p.addnext(el)
    return Paragraph(el, ref_p._parent)


def new_para_after(ref_p, segments, fs=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                   bold=False):
    np_ = para_after(ref_p)
    np_.alignment = align
    for text, hl in segments:
        r = _style(np_.add_run(text), fs=fs, hl=hl)
        r.bold = bold
    return np_


def set_cell(cell, text, hl=True, fs=9, bold=False):
    for pp in cell.paragraphs:
        for r in list(pp.runs):
            r._element.getparent().remove(r._element)
    r = _style(cell.paragraphs[0].add_run(str(text)), fs=fs, hl=hl)
    r.bold = bold
    cell.paragraphs[0].paragraph_format.space_after = Pt(2)


def table_by_caption(prefix):
    """Resolve a table by the caption paragraph that precedes it.

    doc.tables indices shift as soon as a table is inserted, so positional
    lookup silently targets the wrong table. Caption lookup is stable.
    """
    from docx.table import Table as _T
    from docx.text.paragraph import Paragraph as _P
    kids = list(doc.element.body.iterchildren())
    hit = None
    for i, ch in enumerate(kids):
        if ch.tag.endswith("}p") and _P(ch, doc).text.strip().startswith(prefix):
            hit = i
            break
    if hit is None:
        raise KeyError(prefix)
    for ch in kids[hit + 1:]:
        if ch.tag.endswith("}tbl"):
            return _T(ch, doc)
    raise KeyError(f"no table after {prefix}")


def para_after_table(tbl):
    """Insert a paragraph directly AFTER a table element."""
    el = OxmlElement("w:p")
    tbl._tbl.addnext(el)
    return Paragraph(el, tbl._parent)


def build_table_after(ref_p, header, rows, fs=9, hl_all=True):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        set_cell(t.rows[0].cells[i], h, hl=hl_all, fs=fs, bold=True)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            set_cell(cells[i], v, hl=hl_all, fs=fs)
    ref_p._p.addnext(t._tbl)
    return t


# ═══════════════════════════════════════════════ 1. EQUATIONS (LaTeX)
EQS = [
 (41, r"D=\left\{(\mathbf{x}_{i,t},\mathbf{s}_{i},y_{i,t})\;:\;i=1,\dots,N;\;t=1,\dots,T\right\},\quad N=1068,\;T=252", "1"),
 (43, r"\mu_{m}^{(r)}=\frac{1}{|S_{m}|}\sum_{(i,t)\in S_{m}}r_{i,t},\quad S_{m}=\left\{(i,t):\mathrm{month}(t)=m,\;t\in\mathrm{TRAIN}\right\}", "2a"),
 (44, r"a_{i,t}^{(r)}=\mathrm{clip}\!\left(\frac{r_{i,t}-\mu_{\mathrm{month}(t)}^{(r)}}{\sigma^{(r)}_{\mathrm{TRAIN}}},\,-5,\,+5\right)", "2b"),
 (46, r"y_{i,e}^{(H)}=\frac{1}{H}\sum_{k=0}^{H-1}a_{i,\,e+k}^{(r)},\quad H=3", "3"),
 (63, r"E_{\mathrm{split}}=\left\{e\;:\;e\geq t_{\min}(\mathrm{split}),\;\;e+H-1\leq t_{\max}(\mathrm{split}),\;\;e-L\geq 1\right\},\quad L=12", "4"),
 (65, r"\tilde{x}_{i,t,j}=\frac{x_{i,t,j}-\mu_{j}^{\mathrm{TRAIN}}}{\sigma_{j}^{\mathrm{TRAIN}}},\quad \sigma_{j}^{\mathrm{TRAIN}}>0", "5"),
 (68, r"\mathbf{x}_{i,t}=\left[\mathbf{x}^{\mathrm{state}}_{i,t}\;\Vert\;\mathbf{x}^{\mathrm{force}}_{i,t}\;\Vert\;\mathbf{x}^{\mathrm{hum}}_{i,t}\;\Vert\;\mathbf{x}^{\mathrm{cal}}_{t}\right]\in\mathbb{R}^{32},\quad 5+22+2+3=32", "6"),
 (71, r"\mathbf{s}_{i}=\left[\mathrm{elev}_{\mu},\,\mathrm{elev}_{\sigma},\,\mathrm{relief},\,\mathrm{slope}_{\mu},\,\mathrm{slope}_{\sigma},\,\mathrm{rough},\,\mathrm{karst},\,\mathrm{area}\right]\in\mathbb{R}^{8}", "7"),
 (73, r"\mathbf{x}^{\mathrm{ann}}_{i,k}=\left[\bar{\mathbf{x}}_{i,k}\;\Vert\;\bar{\mathbf{x}}_{i,k}-\bar{\mathbf{x}}_{i,k-1}\right]\in\mathbb{R}^{64},\qquad \bar{\mathbf{x}}_{i,k}=\frac{1}{12}\sum_{t\in\mathrm{year}\,k}\tilde{\mathbf{x}}_{i,t}", "8"),
 (75, r"\mathbf{h}_{i,e}=\mathrm{LN}\!\left(\mathrm{GRU}_{m}(\tilde{\mathbf{x}}_{i,\,e-L:e-1})\right)+a_{\mathrm{year}}\,\mathrm{LN}\!\left(\mathrm{GRU}_{a}(\mathbf{x}^{\mathrm{ann}}_{i})\right)+a_{\mathrm{state}}\,f_{\mathrm{state}}(\tilde{\mathbf{x}}_{i,\,e-1})", "9"),
 (83, r"\mathbf{h}_{i}\leftarrow\mathbf{h}_{i}+a_{\mathrm{sp}}\,\mathrm{LN}\!\left(\mathrm{Attn}(\mathbf{h},\mathbf{A}_{\mathrm{sp}})\right)_{i}+a_{\mathrm{hy}}\,\mathrm{LN}\!\left(\mathrm{Attn}(\mathbf{h},\mathbf{A}_{\mathrm{hy}})\right)_{i}", "10"),
 (86, r"\mathbf{w}_{i}=\mathrm{softmax}\!\left(\mathbf{W}_{g}\mathbf{s}_{i}\right),\qquad \theta_{i}=\sum_{k=1}^{4}w_{i,k}\,\mathrm{Expert}_{k}(\mathbf{z}_{i})", "11"),
 (88, r"\rho_{i}=\mathrm{sigm}\!\left(\theta^{(1)}_{i}\right),\quad \varsigma_{i}=\mathrm{sigm}\!\left(\theta^{(2)}_{i}\right),\quad \kappa_{i}=\tanh\!\left(\theta^{(3)}_{i}\right),\quad \delta_{i}=\theta^{(4)}_{i}", "12a"),
 (89, r"\hat{y}_{i,e}=\rho_{i}\,y^{(H)}_{i,\,e-H}+\varsigma_{i}\,y^{(H)}_{i,\,e-12}+\kappa_{i}\,f_{\mathrm{eff},i}+\delta_{i}+a_{\mathrm{res}}\,\mathrm{direct}(\mathbf{z}_{i})", "12b"),
 (96, r"L_{\mathrm{pred}}=\frac{\sum_{i,e}m_{i,e}\left\Vert\hat{\mathbf{y}}_{i,e}-\mathbf{y}_{i,e}\right\Vert^{2}}{d_{\mathrm{out}}\sum_{i,e}m_{i,e}}", "13"),
 (98, r"L_{\mathrm{rec}}=\left\Vert g_{\mathrm{rec}}(\mathbf{h})-\tilde{\mathbf{x}}_{\cdot,\,e-1}\right\Vert^{2},\qquad L_{\mathrm{smooth}}=\mathrm{mean}\!\left[\left(\rho_{\cdot,\,e+1}-\rho_{\cdot,e}\right)^{2}\right]", "14a"),
 (99, r"L_{\mathrm{lap}}=\frac{1}{B\,d_{\mathrm{out}}}\sum_{b=1}^{B}\sum_{o=1}^{d_{\mathrm{out}}}\mathbf{r}_{b,o}^{\mathsf{T}}\,\mathbf{L}\,\mathbf{r}_{b,o},\qquad \mathbf{r}=\left(\hat{\mathbf{y}}-\mathbf{y}\right)\odot\mathbf{m}", "14b"),
 (101, r"L_{\mathrm{bal}}=\sum_{k=1}^{4}u_{k}\log u_{k},\qquad \mathbf{u}=\frac{1}{N}\sum_{i=1}^{N}\mathbf{w}_{i}", "15a"),
 (102, r"L_{\mathrm{asym}}=\mathrm{mean}\!\left[\max\!\left(\left|\hat{y}\right|-5,\,0\right)^{2}\right]", "15b"),
 (112, r"\mathrm{RMSE}=\sqrt{\frac{1}{n}\sum_{i,e}\left(\hat{y}_{i,e}-y_{i,e}\right)^{2}}", "16a"),
 (113, r"\mathrm{MAE}=\frac{1}{n}\sum_{i,e}\left|\hat{y}_{i,e}-y_{i,e}\right|", "16b"),
 (115, r"R^{2}=1-\frac{\sum_{i,e}\left(\hat{y}_{i,e}-y_{i,e}\right)^{2}}{\sum_{i,e}\left(y_{i,e}-\bar{y}\right)^{2}}", "17"),
 (117, r"R^{2}_{\mathrm{within}}=1-\frac{\sum_{i,e}\left[\left(\hat{y}_{i,e}-\bar{\hat{y}}_{i}\right)-\left(y_{i,e}-\bar{y}_{i}\right)\right]^{2}}{\sum_{i,e}\left(y_{i,e}-\bar{y}_{i}\right)^{2}},\qquad \bar{y}_{i}=\frac{1}{|E|}\sum_{e\in E}y_{i,e}", "18"),
 (119, r"I=\frac{n}{S_{0}}\cdot\frac{\sum_{i}\sum_{j}W_{ij}\,r_{i}\,r_{j}}{\sum_{i}r_{i}^{2}},\qquad S_{0}=\sum_{i}\sum_{j}W_{ij}", "19"),
 (121, r"d=1-\frac{\sum_{i,e}\left(y_{i,e}-\hat{y}_{i,e}\right)^{2}}{\sum_{i,e}\left(\left|\hat{y}_{i,e}-\bar{y}\right|+\left|y_{i,e}-\bar{y}\right|\right)^{2}},\qquad \mathrm{KGE}=1-\sqrt{(r-1)^{2}+\left(\frac{\sigma_{p}}{\sigma_{o}}-1\right)^{2}+\left(\frac{\mu_{p}-\mu_{o}}{\sigma_{o}}\right)^{2}}", "20"),
]
for idx, latex, num in EQS:
    set_eq(idx, latex, num)

# total-loss equation is missing entirely -> add as (15c) after (15b)
new_para_after(P[102], [(r"L=L_{\mathrm{pred}}+0.005\,L_{\mathrm{rec}}+0.002\,L_{\mathrm{smooth}}"
                        r"+0.15\,L_{\mathrm{lap}}+0.001\,L_{\mathrm{asym}}+0.001\,L_{\mathrm{bal}}"
                        "\t(15c)", True)],
               fs=11, align=WD_ALIGN_PARAGRAPH.LEFT)
LOG.append("added total-loss equation (15c)")

print(f"equations inserted: {len(EQS)} + 1 added")


# ═════════════════════════════ 2. DIMENSION MISMATCH (Eq 6 text, para 69)
rewrite(69, [
 ("where ", False),
 ("x_(i,t) is the dynamic observation vector of county i in month t and the double bar "
  "denotes concatenation. The four blocks sum to the stated 32 channels. The state block "
  "holds five channels: the raw land surface temperature, kNDVI and NDVI of the county "
  "together with the deseasonalised land-surface-temperature and kNDVI anomalies. The "
  "forcing block holds twenty-two channels: thirteen TerraClimate variables "
  "(precipitation, reference and actual evapotranspiration, climatic water deficit, "
  "runoff, maximum and minimum temperature, vapour pressure deficit, soil moisture, "
  "downward shortwave radiation, the Palmer Drought Severity Index, snow water equivalent "
  "and climatic water balance), eight standardised anomalies derived from them (the heat "
  "and drought anomalies together with the standardised forms of maximum temperature, "
  "precipitation, soil moisture, radiation, vapour pressure deficit and drought severity) "
  "and the relative clear-sky observation count. The human block holds two channels, the "
  "mean and total nighttime-light radiance of the county. The calendar block holds three "
  "channels, the sine and cosine of the month index and a linear year fraction, so that "
  "December and January remain adjacent. In the response decoder the five state channels "
  "are treated as state and the twenty-two forcing channels as forcing.", True)])

# para 34: "Twelve climate variables" -> thirteen, and the DataV hyperlink -> citation
rewrite(34, [
 ("The dataset contains 269,136 rows and 58 columns of monthly observations across the "
  "upper, middle and lower reaches. It includes 1,068 county time series, each covering "
  "252 consecutive months from January 2000 to December 2020. Each series is identified by "
  "its GB/T 2260 administrative code. Land surface temperature is obtained from MOD11A2 "
  "[51]. NDVI is obtained from MOD13Q1 [52], and kNDVI is computed from it following "
  "Camps-Valls et al. (2021) [46]. ", False),
 ("Thirteen climate variables are taken from TerraClimate [45], from which a further "
  "eight standardised anomalies are derived.", True),
 (" Elevation, relief, slope and roughness are derived from the Copernicus GLO-30 DEM "
  "[53]. Carbonate-rock fraction is taken from the World Karst Aquifer Map [50]. ", False),
 ("Nighttime light radiance is obtained from the harmonised DMSP-VIIRS series of Li et "
  "al. (2020) [54], using the extended release of that record that covers 2000-2020 "
  "[58]; the original publication documents 1992-2018 only.", True),
 (" All variables are aggregated to county boundaries from DataV.GeoAtlas ", False),
 ("[57]", True),
 (". Counties that were merged or renamed during the study period are mapped to their "
  "2020 boundaries, so that each series refers to a consistent spatial unit. The panel "
  "also includes clear-sky observation counts, standardised heat and drought anomalies, "
  "contiguity adjacency and directed flow links ordered by elevation.", False)])

# ═════════════════════════════ 3. YREB 40% STATISTIC NEEDS A SOURCE (para 17)
rewrite(17, [
 ("Ecological resilience is the capacity of an ecosystem to resist disturbance and return "
  "to its previous functional state. It has become central to environmental governance in "
  "rapidly developing regions. The Yangtze River Economic Belt (YREB) supports more than "
  "40% of China's population and economic output", False),
 (" [2], [38]", True),
 (". It spans subtropical lowlands, karst uplands and a dense river network. In this "
  "region, resilience assessment directly informs ecological redlining, restoration "
  "investment and basin-wide coordination. A reliable assessment must meet three "
  "requirements. It must separate resistance from recovery. It must capture the seasonal "
  "cycles of disturbance and recovery through which resilience operates. It must also "
  "represent the spatial and hydrological pathways along which disturbances spread. "
  "Whether current methods meet these requirements decides if resilience estimates can "
  "guide future management or only describe past conditions.", False)])

# ═════════════════════════════ 4. 30% EDGE CUT + 8x UPDATES (para 84)
NB = 42        # ceil(166/4) window batches per epoch
EXP = NB * (3 + 0.25)
V1 = 21        # ceil(166/8) with the un-chunked batch size
rewrite(84, [
 ("Because the full graph of 1,068 counties yields only one sample per target month, the "
  "belt is divided into three spatially contiguous partitions of 356 counties each. ", False),
 ("Each partition retains between 29.9% and 31.3% of the contiguity edges, so any single "
  "partition step sees roughly one third of the graph; the 8.0% of edges that join "
  "different partitions are absent from every partition step and are seen only on "
  "full-graph steps. Training alternates between partition steps and full-graph steps, "
  "the latter taken with probability 0.25 after each window batch. With 166 training "
  f"windows in batches of four, this raises the number of optimiser updates per epoch "
  f"from 21 under un-chunked training to {EXP:.0f} in expectation, a factor of "
  f"{EXP/V1:.1f}, and to 168 in the limiting case where every batch also receives a "
  "full-graph step. Partitioning therefore both increases the update budget and acts as a "
  "form of edge dropout.", True),
 (" The propagated representation is concatenated with the encoded static context to form "
  "the joint representation z used by all heads.", False)])

# ═════════════════════════════ 5. DECODER INITIALISATION WORDING (after para 90)
new_para_after(P[90], [
 ("The coefficient heads are initialised so that training begins from a "
  "seasonally-weighted blend rather than from a random mapping: the bias of the "
  "seasonal-carry unit is set to 0.8 and that of the persistence unit to 0.0, which after "
  "the logistic transform gives an initial seasonal-carry coefficient of 0.69 against a "
  "persistence coefficient of 0.50, with the forcing sensitivity and offset initialised at "
  "zero. This is an initialisation toward the seasonal-naive solution and not a "
  "reproduction of it, since the exact seasonal-naive predictor requires a seasonal-carry "
  "coefficient of one and a persistence coefficient of zero. No parameters are transferred "
  "from a previously trained model, so the procedure is an initialisation rather than "
  "pre-training.", True)])
LOG.append("added decoder-initialisation clarification after para 90")

# ═════════════════════════════ 6. TABLE 3: H=12 values + stride-3 column
t3 = table_by_caption("Table 3.")
R12 = RIDGE.loc[RIDGE.H == 12].iloc[0]
set_cell(t3.rows[6].cells[4], f"{R12.Ridge:.4f}", fs=9)
set_cell(t3.rows[6].cells[5], f"{R12.ridge_within_r2:.4f}", fs=9)
t3.add_column(Inches(0.85))
set_cell(t3.rows[0].cells[7], "Ridge stride-H R2", fs=9, bold=True)
for ri in range(1, 7):
    h = int(float(t3.rows[ri].cells[0].text.split()[0]))
    w = WSCAN.loc[WSCAN.H == h]
    set_cell(t3.rows[ri].cells[7],
             f"{w.stride_r2.iloc[0]:.4f}" if len(w) else "not estimable", fs=9)
LOG.append("Table 3: H=12 ridge values filled, stride-H column added")

# para 50: replace the "cannot be estimated reliably" claim with the measured value
rewrite(50, [
 ("Table 3 shows that pooled skill rises with the horizon, but so does the share of "
  "variance that is purely cross-sectional. At H = 12 the seasonal-naive predictor already "
  "reaches R2 = 0.9346, and a learned model adds almost nothing. ", False),
 (f"The within-county ceiling at that horizon is strongly negative, "
  f"{R12.ridge_within_r2:.4f}, because a twelve-month mean barely changes within a "
  "three-year test period, so the estimator cannot track the little temporal variation "
  "that remains and performs worse than each county's own mean.", True),
 (" Within-county skill peaks between H = 3 and H = 4. Moving from H = 3 to H = 4 raises "
  "within-county R2 by only 0.007, while the between-county share increases from 65.2% to "
  "70.0%. We therefore selected H = 3. It matches one climatological season and exceeds "
  "the pooled threshold of 0.82, which we fixed before model development as the minimum "
  "acceptable pooled skill.", False),
 (" That threshold is stated here where the audit is introduced and is applied again to "
  "the final model in Section 4.6.", True),
 (" It also stays close to the maximum temporal skill while limiting the variance that is "
  "trivially predictable from county identity. Figure 2 shows how predictability and the "
  "trivially predictable component change with the horizon.", False)])

print("text corrections + Table 3 done")


# ═════════════════════════════ 7. TABLE 5 (missing) — parameter distribution
PARAMS = [("Monthly temporal encoder (N5)", 93504), ("Annual temporal encoder (N5)", 46849),
          ("Forcing encoder", 11520), ("State encoder", 9889),
          ("Spatial graph attention (N2b)", 37441),
          ("Hydrological graph attention (N2a)", 37441),
          ("Static context encoder", 864), ("Lithology expert gate (N3)", 36),
          ("Attribution head (N6)", 772), ("Coefficient experts (N1, four experts)", 77216),
          ("Forcing-effect head (N1)", 18722), ("Free residual head", 18723),
          ("Reconstruction head (N4)", 3104)]
TOTP = sum(v for _, v in PARAMS)
assert TOTP == 356081, TOTP
BASE = [("DRSEI (AE + LSTM)", 142274), ("STGCN", 101730),
        ("Temporal Fusion Transformer", 1510850)]

cap5 = new_para_after(P[93], [
 ("Table 5. Parameter distribution of the Proposed PERSIST and of the retrained deep "
  "baselines", True)], bold=True)
rows5 = [[k, f"{v:,}", f"{100*v/TOTP:.2f} %"] for k, v in PARAMS]
rows5.append(["PERSIST total", f"{TOTP:,}", "100.00 %"])
for k, v in BASE:
    rows5.append([f"Baseline: {k}", f"{v:,}", f"{v/TOTP:.2f}x PERSIST"])
tbl5 = build_table_after(cap5, ["Component", "Parameters", "Share"], rows5, fs=10)
tail5 = para_after_table(tbl5)
_style(tail5.add_run(
    f"PERSIST contains {TOTP:,} trainable parameters. The temporal encoders and the "
    f"coefficient experts together account for "
    f"{100*(93504+46849+77216)/TOTP:.1f}% of them, while the gating and attribution "
    f"pathways that supply lithological and stream conditioning require fewer than 1,000 "
    f"parameters combined. The Temporal Fusion Transformer baseline carries "
    f"{1510850:,} parameters, which is the {1510850/TOTP:.2f}x ratio quoted in Section "
    f"4.6."), fs=12, hl=True)
tail5.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
LOG.append("Table 5 inserted with per-component and baseline parameter counts")

# ═════════════════════════════ 8. kNDVI RESULTS (Table 8 index 6 + para 130)
t8 = table_by_caption("Table 8.")
KN = [("kNDVI RMSE (secondary target)", f"{M['kndvi_ds_RMSE']:.4f}"),
      ("kNDVI MAE (secondary target)", f"{M['kndvi_ds_MAE']:.4f}"),
      ("kNDVI R2 (secondary target)", f"{M['kndvi_ds_R2']:.4f}"),
      ("kNDVI Pearson correlation", f"{M['kndvi_ds_PearsonR']:.4f}"),
      ("kNDVI Willmott index", f"{M['kndvi_ds_WillmottD']:.4f}"),
      ("kNDVI KGE", f"{M['kndvi_ds_KGE']:.4f}"),
      ("kNDVI bias", f"{M['kndvi_ds_Bias']:+.4f}"),
      ("R2 in native kNDVI units", f"{M['raw_kndvi_R2']:.4f}"),
      ("LST-only RMSE / R2 (primary target)",
       f"{M['lst_ds_RMSE']:.4f} / {M['lst_ds_R2']:.4f}")]
for lab, val in KN:
    c = t8.add_row().cells
    set_cell(c[0], lab, fs=12)
    set_cell(c[1], val, fs=12)
LOG.append("Table 8: nine kNDVI / per-target rows added")

new_para_after(P[130], [
 (f"Both declared targets are reported. The pooled figures above average the primary "
  f"land-surface-temperature anomaly and the secondary kNDVI anomaly. Taken separately, "
  f"the thermal target attains R2 = {M['lst_ds_R2']:.4f} with an RMSE of "
  f"{M['lst_ds_RMSE']:.4f}, whereas the vegetation target attains R2 = "
  f"{M['kndvi_ds_R2']:.4f} with an RMSE of {M['kndvi_ds_RMSE']:.4f}. kNDVI is the harder "
  f"of the two by a clear margin, which is expected because vegetation anomalies respond "
  f"to management, phenology and land-cover change that the climate forcing does not "
  f"capture, whereas thermal anomalies track the forcing more directly. In native units "
  f"the model explains {100*M['raw_kndvi_R2']:.1f}% of three-month mean kNDVI variance. "
  f"Because the two targets differ this much, the pooled coefficient sits between them and "
  f"should not be compared with single-target results from other studies without stating "
  f"the basis; Section 4.6 returns to this point.", True)])
LOG.append("kNDVI discussion paragraph added after para 130")

# ═════════════════════════════ 9. RIDGE BASIS (para 157)
R3 = RIDGE.loc[RIDGE.H == 3].iloc[0]
W3 = WSCAN.loc[WSCAN.H == 3].iloc[0]
rewrite(157, [
 ("The comparison also exposes a clear weakness. STGCN reaches a residual Moran's I of "
  "0.5473, compared with 0.8172 for PERSIST. Plain graph convolution therefore absorbs "
  "spatial structure better than the gated graph attention used here, even though PERSIST "
  "uses two graphs and is explicitly penalised for spatially correlated residuals. "
  "Together with the ablation result that the graph pathways fall within seed noise, this "
  "indicates that the spatial component of PERSIST needs redesign. A second point concerns "
  "the ridge ceiling in Table 3. ", False),
 (f"The pooled R2 of PERSIST ({M['all_R2']:.4f}) is lower than the ridge estimate "
  f"({R3.Ridge:.4f}), but the two numbers are not computed on the same basis. The ridge "
  f"estimator predicts only the primary land-surface-temperature anomaly, whereas the "
  f"pooled R2 of PERSIST averages both declared targets, including the harder kNDVI "
  f"anomaly at {M['kndvi_ds_R2']:.4f}. Restricted to the same single target and the same "
  f"34 test windows, PERSIST attains R2 = {M['lst_ds_R2']:.4f} against {R3.Ridge:.4f} and "
  f"an RMSE of {M['lst_ds_RMSE']:.4f} against {R3.ridge_rmse:.4f}, so it leads the linear "
  f"ceiling by {M['lst_ds_R2']-R3.Ridge:.4f} in R2 and by "
  f"{100*(R3.ridge_rmse-M['lst_ds_RMSE'])/R3.ridge_rmse:.1f}% in RMSE. The within-county "
  f"and stride-3 coefficients quoted throughout are already single-target and need no "
  f"adjustment: {M['within_county_R2']:.4f} against {R3.ridge_within_r2:.4f} and "
  f"{M['strideH_R2']:.4f} against {W3.stride_r2:.4f}, the latter now also reported for "
  f"every horizon in Table 3. PERSIST therefore exceeds the linear ceiling on every metric "
  f"once the comparison is placed on a common footing. The two-target pooled coefficient "
  f"remains the least discriminating quantity reported here because it mixes an easier "
  f"thermal target with a harder vegetation target, which is why within-county R2 is "
  f"treated as the primary measure. The final model also clears the pooled acceptance "
  f"threshold of 0.82 fixed in Section 3.2.", True),
 (" Figure 6 compares all models on RMSE, pooled R2, within-county R2 and residual "
  "Moran's I.", False)])

print("Table 5, kNDVI, ridge basis done")


# ═════════════════════════════ 10. SIGNIFICANCE TESTS (Table 11 index 9 + text)
def sg(label, var):
    d = SIG[f"{label}|{var}"]
    star = ("p < 0.001" if d["p"] < .001 else
            f"p = {d['p']:.3f}" if d["p"] >= .001 else "")
    return d, star


t11 = table_by_caption("Table 11.")
t11.add_column(Inches(1.15))
set_cell(t11.rows[0].cells[6], "Omnibus test vs stratum", fs=9, bold=True)
rr, _ = sg("Reach", "resistance")
rk, _ = sg("Lithology", "resistance")
REACH_NOTE = (f"Kruskal-Wallis across reaches: resistance H = {rr['H']:.1f}, p < 0.001, "
              f"eps2 = {rr['eps2']:.3f}")
KARST_NOTE = (f"Kruskal-Wallis across lithological classes: resistance H = {rk['H']:.1f}, "
              f"p < 0.001, eps2 = {rk['eps2']:.3f}")
for ri in range(1, len(t11.rows)):
    lab = t11.rows[ri].cells[0].text
    set_cell(t11.rows[ri].cells[6],
             REACH_NOTE if lab.startswith("Reach") else
             (KARST_NOTE if lab.startswith("Lithology") else "-"), fs=8)

cap12 = new_para_after(P[165], [
 ("Table 12. Significance of the reach and lithology contrasts in the recovered resilience "
  "parameters. Tests use county-level means (n = 1,068) because county-window pairs are "
  "not independent within a county. Pairwise p-values are Holm-corrected [59].", True)],
 bold=True)
rows12 = []
for label, group in [("Reach", "Reach"), ("Lithology", "Lithology")]:
    for var, nice in [("resistance", "Resistance"), ("recovery", "Recovery rate"),
                      ("sig", "Seasonal carry")]:
        d = SIG[f"{group}|{var}"]
        pv = "< 0.001" if d["p"] < .001 else f"{d['p']:.3f}"
        verdict = ("significant" if d["p"] < .05 else "not significant")
        worst = max(d["pairs"], key=lambda x: x[3])
        rows12.append([label, nice, f"{d['H']:.1f}", pv, f"{d['eps2']:.3f}",
                       f"{worst[3]:.3f}" if worst[3] >= .001 else "< 0.001", verdict])
tbl12 = build_table_after(cap12, ["Stratum", "Parameter", "Kruskal-Wallis H(2)", "p",
                                  "Epsilon-squared", "Largest pairwise p (Holm)",
                                  "Verdict"], rows12, fs=9)
tailsig = para_after_table(tbl12)
_style(tailsig.add_run(
    f"Kruskal-Wallis tests on county-level means confirm the asymmetry described above "
    f"rather than merely restating it. "
    f"Resistance differs very strongly across reaches (H = {rr['H']:.1f}, p < 0.001, "
    f"epsilon-squared = {rr['eps2']:.3f}) and across lithological classes "
    f"(H = {rk['H']:.1f}, p < 0.001, epsilon-squared = {rk['eps2']:.3f}), with every "
    f"pairwise contrast significant after Holm correction; the upstream-versus-downstream "
    f"contrast in resistance has a Cliff delta of +0.770 and the non-karst-versus-"
    f"majority-karst contrast +0.618, both large effects. Recovery rate, by contrast, "
    f"does not differ significantly across reaches "
    f"(H = {SIG['Reach|recovery']['H']:.2f}, p = {SIG['Reach|recovery']['p']:.2f}) or "
    f"across lithological classes (H = {SIG['Lithology|recovery']['H']:.2f}, "
    f"p = {SIG['Lithology|recovery']['p']:.2f}), and no pairwise contrast survives "
    f"correction. Seasonal carry differs significantly in both stratifications "
    f"(reaches H = {SIG['Reach|sig']['H']:.1f}, epsilon-squared = "
    f"{SIG['Reach|sig']['eps2']:.3f}; lithology H = {SIG['Lithology|sig']['H']:.1f}, "
    f"epsilon-squared = {SIG['Lithology|sig']['eps2']:.3f}). The claim that the belt shows "
    f"a spatial pattern in resistance but not in recovery is therefore supported by formal "
    f"testing and not only by the group means. All correlations reported above are "
    f"significant at p < 0.001 except carbonate fraction against recovery rate "
    f"(r = +0.005, p = 0.87), which is consistent with the same conclusion."),
    fs=12, hl=True)
tailsig.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
LOG.append("Table 12 (significance tests) + discussion added")

# ═════════════════════════════ 11. FIGURE 5 CAPTION (two panels, para 147)
rewrite(147, [
 ("Figure 5. ", False),
 ("(a) Change in test RMSE relative to the full PERSIST model for each single-component "
  "ablation, with the shaded band marking +/-2 seed standard deviations; (b) test RMSE of "
  "the full model and of every ablation as seed mean with standard-deviation error bars.",
  True)])

# ═════════════════════════════ 12. TABLE 1 REBUILD (benchmarked + discussed)
t1 = table_by_caption("Table 1.")
for row in list(t1.rows)[1:]:
    row._tr.getparent().remove(row._tr)
T1ROWS = [
 ["Gong et al. (2025) [3]", "Autoencoder + LSTM deep RSEI (DRSEI). Benchmarked in this study.",
  "Learned encoding removes subjective indicator weighting; strongest deep baseline here "
  f"(RMSE {0.3871:.4f}, R2 {0.8314:.4f}).",
  "Per-county and purely temporal: no spatial graph, no hydrological connectivity and no "
  "resilience parameters recovered."],
 ["Yu et al. (2018) [43]", "Spatio-temporal graph convolutional network (STGCN). Benchmarked in this study.",
  f"Best residual Moran's I of all models ({0.5473:.4f}), so plain graph convolution "
  f"absorbs spatial structure most effectively.",
  "Generic regression objective; lowest pooled and within-county skill of the three deep "
  "baselines and no interpretable coefficients."],
 ["Lim et al. (2021) [44]", "Temporal Fusion Transformer (TFT). Benchmarked in this study.",
  f"Highest-capacity alternative for this data shape (1,510,850 parameters; RMSE "
  f"{0.3970:.4f}).",
  "4.24x more parameters than PERSIST for lower skill on every metric; attention weights "
  "are not bounded ecological quantities."],
 ["Li et al. (2026) [2]", "Multi-scale composite resilience index for the YREB.",
  "Establishes reach-level resilience gradients and significant spatial dependence among "
  "counties; independent reference for the present reach contrast.",
  "Descriptive; spatial autocorrelation is diagnosed but never embedded in an estimator, "
  "and resistance is not separated from recovery."],
 ["Yang et al. (2025) [24]", "Karst-specific remote sensing ecological index (KRSEI).",
  "Demonstrates that applying a standard index to carbonate terrain produces systematic "
  "bias; motivates the lithology conditioning used here.",
  "Builds a separate index for karst regions rather than one model conditioned on "
  "lithology, preventing unified analysis of mixed terrain."],
 ["Zhu et al. (2025) [27]", "Multi-method analysis of LST and kNDVI dynamics across the YREB.",
  "Provides the belt-wide trends for the two state variables used here and the independent "
  "check on the extracted trends.",
  "Diagnostic and correlational; no predictive model and no resilience parameterisation."],
 ["Ke et al. (2026) [29]", "Cellular automata coupled with graph attention and a transformer.",
  "Confirms that graph attention with temporal attention outperforms conventional cellular "
  "automata; closest architecture in spirit to PERSIST.",
  "Target is discrete land-use transition rather than a continuous ecological anomaly; no "
  "bounded dynamical coefficients."],
 ["Soula et al. (2026) [28]", "BiLSTM-CNN for large-scale NDVI prediction.",
  "Shows hybrid recurrent-convolutional models can forecast vegetation state over broad "
  "areas.",
  "Convolution over a regular grid cannot represent irregular counties or hydrological "
  "flow; point forecasts only."],
 ["Ridge, AR-12 + climate (this study)",
  "Linear autoregressive reference fitted for the horizon audit and reported in Tables 3 "
  "and 10.",
  f"Establishes the predictability ceiling: single-target R2 {R3.Ridge:.4f}, within-county "
  f"{R3.ridge_within_r2:.4f}, stride-3 {W3.stride_r2:.4f}.",
  "Linear and per-county; no spatial structure, no lithological conditioning and no "
  "bounded resilience coefficients."],
 ["PERSIST (proposed)",
  "Process-informed spatiotemporal decoder with dual-graph message passing and "
  "lithology-adaptive experts.",
  f"Single-target R2 {M['lst_ds_R2']:.4f} and within-county R2 {M['within_county_R2']:.4f}, "
  f"exceeding every reference; recovers bounded recovery rate and resistance per county.",
  "Residual Moran's I (0.8172) worse than STGCN; five of seven components fall within seed "
  "noise; hydrological graph is a terrain proxy."],
]
for row in T1ROWS:
    cells = t1.add_row().cells
    for i, v in enumerate(row):
        set_cell(cells[i], v, fs=8)
rewrite(25, [
 ("Table 1. ", False),
 ("Comparison of the proposed PERSIST with the methods benchmarked in this study and with "
  "the principal methods discussed in Section 2.", True)])
LOG.append("Table 1 rebuilt around benchmarked and discussed methods")

print("significance, Fig 5 caption, Table 1 done")


# ═════════════════════════════ 13. REFERENCES: no bare URLs, add missing sources
def find_ref(prefix):
    """Return the live Paragraph object; the cached list P is stale after inserts."""
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    raise KeyError(prefix)


def rewrite_p(p, segments, fs=12):
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    for text, hl in segments:
        _style(p.add_run(text), fs=fs, hl=hl)
    return p


# [45] TerraClimate bare URL -> full citation
rewrite_p(find_ref("[45]"), [
 ("[45] Abatzoglou, J.T., Dobrowski, S.Z., Parks, S.A. and Hegewisch, K.C., 2018. "
  "TerraClimate, a high-resolution global dataset of monthly climate and climatic water "
  "balance from 1958-2015. Scientific Data, 5, 170191. "
  "https://doi.org/10.1038/sdata.2017.191. Accessed via the Microsoft Planetary Computer "
  "STAC collection 'terraclimate'.", True)])

# [50] WOKAM bare URL -> versioned citation
rewrite_p(find_ref("[50]"), [
 ("[50] Goldscheider, N., Chen, Z., Auler, A.S., Bakalowicz, M., Broda, S., Drew, D., "
  "Hartmann, J., Jiang, G., Moosdorf, N., Stevanovic, Z. and Veni, G., 2020. Global "
  "distribution of carbonate rocks and karst water resources. Hydrogeology Journal, 28(5), "
  "pp.1661-1677. https://doi.org/10.1007/s10040-020-02139-y. Carbonate-rock polygons taken "
  "from the World Karst Aquifer Map (WOKAM), 2017 edition, scale 1:40,000,000, BGR, IAH, "
  "KIT and UNESCO, distributed at "
  "https://www.bgr.bund.de/whymap/EN/Maps_Data/Wokam/wokam_node_en.html.", True)])

# [54] note the coverage limit explicitly
p54 = find_ref("[54]")
old54 = p54.text.strip()
rewrite_p(p54, [
 (old54.rstrip(".") if old54.endswith(".") else old54, False),
 (" This publication documents the 1992-2018 record; the 2000-2020 series used here is "
  "taken from the extended release of the same repository, cited separately as [58].",
  True)])

# new references [57]-[59] appended after the last reference paragraph
last = len(doc.paragraphs) - 1
while not doc.paragraphs[last].text.strip().startswith("["):
    last -= 1
anchor = doc.paragraphs[last]
NEW_REFS = [
 "[57] DataV.GeoAtlas, 2023. China administrative boundary service, county level "
 "(GB/T 2260 codes). Alibaba Cloud. "
 "https://datav.aliyun.com/portal/school/atlas/area_selector (accessed 2025).",
 "[58] Li, X., Zhou, Y., Zhao, M. and Zhao, X., 2020. A harmonized global nighttime light "
 "dataset, extended annual release covering 1992-2021. figshare dataset. "
 "https://doi.org/10.6084/m9.figshare.9828827. The 2000-2020 subset of this release is "
 "used in the present study.",
 "[59] Holm, S., 1979. A simple sequentially rejective multiple test procedure. "
 "Scandinavian Journal of Statistics, 6(2), pp.65-70.",
]
for ref in NEW_REFS:
    anchor = new_para_after(anchor, [(ref, True)])
LOG.append("references: [45] and [50] de-URLed, [54] annotated, [57]-[59] added")

# Data Availability: point at the new reference numbers instead of inline links
try:
    ida = find_ref("Data Availability")
except KeyError:
    ida = None
if ida:
    rewrite_p(ida, [
     ("Data Availability: ", False),
     ("All source datasets are publicly available and are cited as numbered references "
      "rather than inline links: MOD11A2 [51], MOD13Q1 [52], TerraClimate [45], Copernicus "
      "GLO-30 [53], the World Karst Aquifer Map [50], the harmonised nighttime-light series "
      "[54] and its extended release [58], and the DataV.GeoAtlas county boundaries [57]. "
      "The derived county panel, the trained model weights, the per-seed metrics and the "
      "recovered resilience coefficients are available from the corresponding author on "
      "request.", True)])

# ═════════════════════════════ save
DST.parent.mkdir(parents=True, exist_ok=True)
doc.save(str(DST))
print("\n".join(f"  - {x}" for x in LOG))
print(f"\nSAVED -> {DST}")
print(f"  paragraphs {len(doc.paragraphs)}  tables {len(doc.tables)}")
